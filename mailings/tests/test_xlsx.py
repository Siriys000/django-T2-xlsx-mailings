from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.exceptions import ValidationError
from django.test import SimpleTestCase
from openpyxl import Workbook

from mailings.models import MailingMessage
from mailings.xlsx import XlsxFormatError, build_mailing_message, iter_xlsx_rows


HEADERS = ["external_id", "user_id", "email", "subject", "message"]


class XlsxReaderTests(SimpleTestCase):
    def setUp(self):
        self.temp_directory = TemporaryDirectory()
        self.addCleanup(self.temp_directory.cleanup)

    def write_workbook(self, rows, filename="mailings.xlsx"):
        path = Path(self.temp_directory.name) / filename
        workbook = Workbook()
        worksheet = workbook.active
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()
        return path

    def test_reads_headers_in_any_order_and_ignores_extra_columns(self):
        path = self.write_workbook(
            [
                [
                    " subject ",
                    "ignored",
                    "external_id",
                    "message",
                    "email",
                    "user_id",
                ],
                [
                    "Subject",
                    "not imported",
                    "mailing-001",
                    "Message",
                    "user@example.com",
                    42,
                ],
            ]
        )

        rows = list(iter_xlsx_rows(path))

        self.assertEqual(
            rows,
            [
                (
                    2,
                    {
                        "external_id": "mailing-001",
                        "user_id": 42,
                        "email": "user@example.com",
                        "subject": "Subject",
                        "message": "Message",
                    },
                )
            ],
        )

    def test_ignores_fully_blank_rows(self):
        path = self.write_workbook(
            [
                HEADERS,
                [None, None, None, None, None],
                ["mailing-001", None, None, None, None],
            ]
        )

        rows = list(iter_xlsx_rows(path))

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0][0], 3)

    def test_rejects_missing_required_headers(self):
        path = self.write_workbook(
            [["external_id", "user_id", "email", "subject"]]
        )

        with self.assertRaisesRegex(XlsxFormatError, "message"):
            list(iter_xlsx_rows(path))

    def test_rejects_empty_active_sheet(self):
        path = self.write_workbook([])

        with self.assertRaisesRegex(XlsxFormatError, "empty"):
            list(iter_xlsx_rows(path))

    def test_rejects_duplicate_required_headers(self):
        path = self.write_workbook(
            [["external_id", "user_id", "email", "subject", "message", "email"]]
        )

        with self.assertRaisesRegex(XlsxFormatError, "email"):
            list(iter_xlsx_rows(path))

    def test_reads_active_sheet(self):
        path = Path(self.temp_directory.name) / "mailings.xlsx"
        workbook = Workbook()
        workbook.active.append(["not", "the", "required", "headers"])
        worksheet = workbook.create_sheet("Mailings")
        worksheet.append(HEADERS)
        worksheet.append(
            ["mailing-001", 42, "user@example.com", "Subject", "Message"]
        )
        workbook.active = 1
        workbook.save(path)
        workbook.close()

        rows = list(iter_xlsx_rows(path))

        self.assertEqual(rows[0][1]["external_id"], "mailing-001")

    def test_rejects_non_xlsx_file(self):
        path = self.write_workbook([HEADERS], filename="mailings.xls")

        with self.assertRaisesRegex(XlsxFormatError, r"\.xlsx"):
            list(iter_xlsx_rows(path))

    def test_rejects_corrupted_xlsx_file(self):
        path = Path(self.temp_directory.name) / "mailings.xlsx"
        path.write_text("not a workbook", encoding="utf-8")

        with self.assertRaises(XlsxFormatError):
            list(iter_xlsx_rows(path))


class MailingRowValidationTests(SimpleTestCase):
    @staticmethod
    def valid_data(**overrides):
        data = {
            "external_id": "mailing-001",
            "user_id": 42,
            "email": "user@example.com",
            "subject": "Subject",
            "message": "Message",
        }
        data.update(overrides)
        return data

    def test_builds_unsaved_message_and_normalizes_scalar_values(self):
        mailing = build_mailing_message(
            self.valid_data(
                external_id=1001.0,
                user_id=" 42 ",
                email=" user@example.com ",
                subject=" Subject ",
                message=" Message body\n",
            )
        )

        self.assertIsInstance(mailing, MailingMessage)
        self.assertIsNone(mailing.pk)
        self.assertEqual(mailing.external_id, "1001")
        self.assertEqual(mailing.user_id, 42)
        self.assertEqual(mailing.email, "user@example.com")
        self.assertEqual(mailing.subject, "Subject")
        self.assertEqual(mailing.message, " Message body\n")

    def test_rejects_invalid_row_values(self):
        invalid_values = [
            ("external_id", True),
            ("external_id", 1.5),
            ("user_id", True),
            ("user_id", 1.5),
            ("user_id", -1),
            ("email", "not-an-email"),
            ("message", "   "),
        ]

        for field, value in invalid_values:
            with self.subTest(field=field, value=value):
                with self.assertRaises(ValidationError) as raised:
                    build_mailing_message(self.valid_data(**{field: value}))

                self.assertIn(field, raised.exception.message_dict)
