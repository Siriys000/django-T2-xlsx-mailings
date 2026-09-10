from collections.abc import Iterator, Mapping, Sequence
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from django.core.exceptions import ValidationError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from mailings.models import MailingMessage

REQUIRED_HEADERS = (
    "external_id",
    "user_id",
    "email",
    "subject",
    "message",
)

WORKBOOK_READ_ERRORS = (
    BadZipFile,
    InvalidFileException,
    OSError,
    ValueError,
    ParseError,
    EOFError,
    KeyError,
)


class XlsxFormatError(ValueError):
    """Raised when a workbook cannot provide the required import structure."""


def iter_xlsx_rows(
    file_path: str | Path,
) -> Iterator[tuple[int, dict[str, object]]]:
    """Yield non-empty rows from the active worksheet without loading it fully."""
    path = Path(file_path)
    if path.suffix.lower() != ".xlsx":
        raise XlsxFormatError("Only .xlsx files are supported.")

    source = None
    workbook = None
    try:
        source = path.open("rb")
        workbook = load_workbook(source, read_only=True, data_only=True)
        rows = workbook.active.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise XlsxFormatError("The active worksheet is empty.") from exc

        header_positions = _get_header_positions(header_row)

        for row_number, values in enumerate(rows, start=2):
            if _is_blank_row(values):
                continue

            yield row_number, {
                header: values[position] if position < len(values) else None
                for header, position in header_positions.items()
            }
    except XlsxFormatError:
        raise
    except WORKBOOK_READ_ERRORS as exc:
        raise XlsxFormatError(f"Cannot read XLSX file: {exc}") from exc
    finally:
        if workbook is not None:
            workbook.close()
        if source is not None:
            source.close()


def build_mailing_message(row_data: Mapping[str, object]) -> MailingMessage:
    """Build and validate an unsaved message from raw worksheet values."""
    mailing = MailingMessage(
        external_id=_normalize_external_id(row_data.get("external_id")),
        user_id=_normalize_user_id(row_data.get("user_id")),
        email=_normalize_text(row_data.get("email"), "email"),
        subject=_normalize_text(row_data.get("subject"), "subject"),
        message=_normalize_text(row_data.get("message"), "message", strip=False),
    )
    mailing.full_clean(validate_unique=False)
    return mailing


def _get_header_positions(header_row: Sequence[object]) -> dict[str, int]:
    positions: dict[str, int] = {}
    duplicates: list[str] = []

    for position, value in enumerate(header_row):
        header = value.strip() if isinstance(value, str) else None
        if header not in REQUIRED_HEADERS:
            continue
        if header in positions:
            duplicates.append(header)
        else:
            positions[header] = position

    if duplicates:
        names = ", ".join(dict.fromkeys(duplicates))
        raise XlsxFormatError(f"Duplicate required headers: {names}.")

    missing = [header for header in REQUIRED_HEADERS if header not in positions]
    if missing:
        raise XlsxFormatError(f"Missing required headers: {', '.join(missing)}.")

    return positions


def _is_blank_row(values: Sequence[object]) -> bool:
    return all(
        value is None or isinstance(value, str) and not value.strip()
        for value in values
    )


def _normalize_external_id(value: object) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        raise ValidationError({"external_id": "Must be a string or integer."})
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if value is None:
        return ""
    raise ValidationError({"external_id": "Must be a string or integer."})


def _normalize_user_id(value: object) -> int | str | None:
    if isinstance(value, bool):
        raise ValidationError({"user_id": "Must be an integer."})
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        normalized = value.strip()
        if normalized.isascii() and normalized.isdecimal():
            return int(normalized)
        if not normalized:
            return ""
    if value is None:
        return None
    raise ValidationError({"user_id": "Must be an integer."})


def _normalize_text(value: object, field: str, *, strip: bool = True) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ValidationError({field: "Must be a string."})

    normalized = value.strip() if strip else value
    if not normalized.strip():
        raise ValidationError({field: "This field cannot be blank."})
    return normalized
